"""Build a PowerPoint slide with the architecture diagram embedded.

Single-slide visual where the SVG-rendered PNG carries its own title.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

HERE = Path(__file__).parent
PNG_PATH = HERE / "architecture.png"
OUT_PATH = HERE / "architecture.pptx"

# 16:9 widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5

# Blank layout
slide = prs.slides.add_slide(prs.slide_layouts[6])

# ---- Background ----
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.line.fill.background()
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0xFA, 0xFB, 0xFD)
spTree = bg._element.getparent()
spTree.remove(bg._element)
spTree.insert(2, bg._element)

# ---- Diagram image (the focus of the slide) ----
img = Image.open(PNG_PATH)
img_w, img_h = img.size
aspect = img_w / img_h

# Reserve a thin footer strip; let the image fill nearly the entire slide
margin_x = 0.4
margin_top = 0.3
footer_strip = 0.55  # space for sources/footer
max_w = SLIDE_W_IN - 2 * margin_x
max_h = SLIDE_H_IN - margin_top - footer_strip

if max_w / aspect <= max_h:
    disp_w = max_w
    disp_h = max_w / aspect
else:
    disp_h = max_h
    disp_w = max_h * aspect

left = (SLIDE_W_IN - disp_w) / 2
top = margin_top + (max_h - disp_h) / 2  # vertically center within available area

slide.shapes.add_picture(
    str(PNG_PATH),
    Inches(left),
    Inches(top),
    width=Inches(disp_w),
    height=Inches(disp_h),
)

# ---- Footer strip with terraform source attribution ----
footer_top = SLIDE_H_IN - footer_strip + 0.05
footer_box = slide.shapes.add_textbox(
    Inches(0.5), Inches(footer_top), Inches(12.333), Inches(0.4)
)
ftf = footer_box.text_frame
ftf.margin_left = ftf.margin_right = ftf.margin_top = ftf.margin_bottom = 0
ftf.word_wrap = True
fp = ftf.paragraphs[0]
fp.alignment = PP_ALIGN.LEFT


def add_run(p, text, color, bold=False, italic=False, size=11):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Calibri"
    r.font.color.rgb = RGBColor(*color)


add_run(fp, "Source: ", (0x1A, 0x22, 0x33), bold=True)
add_run(fp, "terraform/main.tf", (0x3F, 0x4B, 0x5C), italic=True)
add_run(fp, "    │    ", (0xB0, 0xB8, 0xC4))
add_run(fp, "Resources: ", (0x1A, 0x22, 0x33), bold=True)
add_run(fp, "API Gateway REST  •  IAM role/policy  •  SQS main + DLQ (redrive)", (0x3F, 0x4B, 0x5C))

prs.save(OUT_PATH)
print(f"Wrote {OUT_PATH}")
print(f"Image displayed at {disp_w:.2f}\" x {disp_h:.2f}\"  (slide is {SLIDE_W_IN}\" x {SLIDE_H_IN}\")")
